"""Filters/search/sort, search-filter conversion, bulk preview/confirm, CSV/XLSX/JSON export,
full wipe and user deletion."""

from __future__ import annotations

import csv
import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

import pytest
from sqlalchemy import func, select

from app.core.errors import AppError
from app.models import Account, ActionLog, Category, Posting, PreviewToken, Transaction, User
from app.schemas.common import PeriodSpec
from app.schemas.intents import (
    DeleteTransactionIntent,
    ParseContext,
    SearchFilter,
    TargetSpec,
    TransactionChanges,
    UpdateTransactionIntent,
)
from app.services.ledger import actions, bulk, query
from app.services.ledger.common import now
from app.services.ledger.wipe import delete_user, wipe_financial_data
from tests.core.conftest import acc_id, cat_id, idem


async def _seed_ops(api, session, user):  # type: ignore[no-untyped-def]
    food = await cat_id(session, user.id, "Еда")
    taxi = await cat_id(session, user.id, "Такси/транспорт")
    srv = await cat_id(session, user.id, "Серверы/инфра")
    kb = await acc_id(session, user.id, "Кошелёк")
    m = (await api.post("/merchants", json={"name": "vast.ai", "aliases": ["васт"],
                                            "category_id": srv})).json()
    ops = [
        {"kind": "expense", "amount": "120000", "currency": "VND", "category_id": food,
         "description": "обед с командой", "occurred_on": "2026-10-20"},
        {"kind": "expense", "amount": "45000", "currency": "VND", "category_id": taxi,
         "occurred_on": "2026-10-21"},
        {"kind": "expense", "amount": "180", "currency": "USD", "category_id": srv,
         "merchant_id": m["id"], "account_none": True, "occurred_on": "2026-10-22"},
        {"kind": "expense", "amount": "3", "currency": "TON", "category_id": food,
         "account_id": kb, "counterparty": "@bob", "occurred_on": "2026-10-23"},
    ]
    out = []
    for o in ops:
        r = await api.post("/transactions", headers=idem(), json=o)
        assert r.status_code == 201, r.text
        out.append(r.json()["transaction"])
    return out


async def test_filters_search_sort(api, session, user):
    ops = await _seed_ops(api, session, user)
    food = await cat_id(session, user.id, "Еда")

    async def ids(**params):  # type: ignore[no-untyped-def]
        r = await api.get("/transactions", params=params)
        assert r.status_code == 200, r.text
        return [x["id"] for x in r.json()["items"]]

    assert await ids() == [ops[3]["id"], ops[2]["id"], ops[1]["id"], ops[0]["id"]]
    assert await ids(sort="date_asc") == [o["id"] for o in ops]
    assert await ids(date_from="2026-10-21", date_to="2026-10-22") == [ops[2]["id"],
                                                                        ops[1]["id"]]
    assert await ids(category_id=food) == [ops[3]["id"], ops[0]["id"]]
    assert await ids(q="васт") == [ops[2]["id"]]          # merchant alias
    assert await ids(q="КОМАНД") == [ops[0]["id"]]        # description, case-insensitive
    assert await ids(q="такси") == [ops[1]["id"]]          # category name
    assert await ids(q="bob") == [ops[3]["id"]]        # counterparty
    assert await ids(q="100%_") == []                       # LIKE wildcards are escaped
    assert await ids(currency="TON") == [ops[3]["id"]]
    # «покажи траты дороже 100 долларов» → amount_usd
    assert await ids(min_amount="100", amount_currency="USD") == [ops[2]["id"]]
    assert (await ids(sort="amount_desc"))[0] == ops[2]["id"]
    assert await ids(status="needs_account") == []
    assert await ids(is_business=True) == [ops[2]["id"]]
    page = (await api.get("/transactions", params={"limit": 2, "offset": 1})).json()
    assert page["total"] == 4 and len(page["items"]) == 2 and page["offset"] == 1
    assert Decimal(page["totals"]["expense_usd"]) > Decimal("180")
    assert (await api.get("/transactions", params={"limit": 0})).status_code == 422


async def test_search_filter_and_targets(api, session, user):
    ops = await _seed_ops(api, session, user)
    ctx = ParseContext(now=now(), today=date(2026, 10, 23))
    f = query.search_filter_to_filter(
        SearchFilter(period=PeriodSpec(preset="last_7_days"), text="обед", kinds=["expense"],
                     min_amount="100", amount_currency="USD", limit=5),
        ctx)
    assert f.date_from == date(2026, 10, 17) and f.date_to == date(2026, 10, 23)
    assert f.q == "обед" and f.kinds == ["expense"] and f.limit == 5
    assert query.resolve_period(PeriodSpec(preset="this_week"), date(2026, 10, 23)) == (
        date(2026, 10, 19), date(2026, 10, 25))
    assert query.resolve_period(PeriodSpec(preset="last_month"), date(2026, 10, 23)) == (
        date(2026, 9, 1), date(2026, 9, 30))
    assert query.resolve_period(PeriodSpec(preset="yesterday"), date(2026, 10, 1)) == (
        date(2026, 9, 30), date(2026, 9, 30))
    found = await query.resolve_target(
        session, user.id,
        TargetSpec(kind="search", search=SearchFilter(text="такси")), ctx)
    assert [t.id for t in found] == [ops[1]["id"]]
    by_id = await query.resolve_target(
        session, user.id, TargetSpec(kind="id", transaction_id=ops[0]["id"]), ctx)
    assert by_id[0].id == ops[0]["id"]
    # no context → newest created
    last = await query.resolve_target(session, user.id, TargetSpec(kind="last"), ctx)
    assert last[0].id == ops[3]["id"]


async def test_bulk_preview_confirm_and_undo(api, session, user, principal):
    ops = await _seed_ops(api, session, user)
    rows = list(await session.scalars(
        select(Transaction).where(Transaction.id.in_([ops[0]["id"], ops[1]["id"]]))))
    tok = await bulk.create_preview(session, principal, DeleteTransactionIntent(), rows)
    assert tok.summary == "Удалить операций: 2"
    aid, n = await bulk.confirm_preview(session, principal, tok.id)
    assert n == 2
    # repeated confirm → same action (no double delete)
    assert await bulk.confirm_preview(session, principal, tok.id) == (aid, 2)
    live = await session.scalar(select(func.count()).where(
        Transaction.user_id == user.id, Transaction.deleted_at.is_(None)))
    assert live == 2
    act = await session.get(ActionLog, aid)
    assert act.kind == "bulk_delete" and len(act.entity_refs) == 2
    res = await actions.undo_action(session, principal, aid)
    assert res.status == "undone"
    live = await session.scalar(select(func.count()).where(
        Transaction.user_id == user.id, Transaction.deleted_at.is_(None)))
    assert live == 4

    # bulk update with a resolved patch
    taxi = await cat_id(session, user.id, "Такси/транспорт")
    rows = list(await session.scalars(
        select(Transaction).where(Transaction.id.in_([ops[0]["id"], ops[3]["id"]]))
        .execution_options(populate_existing=True)))
    intent = UpdateTransactionIntent(
        target=TargetSpec(kind="search", search=SearchFilter(text="x")),
        changes=TransactionChanges(category_ref="такси"))
    with pytest.raises(AppError):
        await bulk.create_preview(session, principal, intent, rows)  # unresolved ref
    tok = await bulk.create_preview(session, principal, intent, rows,
                                    patch={"category_id": taxi})
    # a record changed meanwhile → stale
    t0 = await session.get(Transaction, ops[0]["id"])
    await api.patch(f"/transactions/{t0.id}", json={"version": t0.version,
                                                    "description": "x"})
    with pytest.raises(AppError) as e:
        await bulk.confirm_preview(session, principal, tok.id)
    assert e.value.code == "preview_stale"
    rows = list(await session.scalars(
        select(Transaction).where(Transaction.id.in_([ops[0]["id"], ops[3]["id"]]))
        .execution_options(populate_existing=True)))
    tok = await bulk.create_preview(session, principal, intent, rows,
                                    patch={"category_id": taxi})
    aid, n = await bulk.confirm_preview(session, principal, tok.id)
    for tid in (ops[0]["id"], ops[3]["id"]):
        assert (await api.get(f"/transactions/{tid}")).json()["category"]["id"] == taxi
    # expired / cancelled
    tok = await bulk.create_preview(session, principal, DeleteTransactionIntent(), rows[:1])
    await bulk.cancel_preview(session, principal, tok.id)
    with pytest.raises(AppError) as e:
        await bulk.confirm_preview(session, principal, tok.id)
    assert e.value.code == "preview_expired"
    rows = list(await session.scalars(select(Transaction).where(Transaction.id == ops[1]["id"])))
    tok = await bulk.create_preview(session, principal, DeleteTransactionIntent(), rows)
    tok_row = await session.get(PreviewToken, tok.id)
    tok_row.expires_at = now() - timedelta(seconds=1)
    await session.flush()
    with pytest.raises(AppError) as e:
        await bulk.confirm_preview(session, principal, tok.id)
    assert e.value.code == "preview_expired"


async def test_export_csv_xlsx_json(api, session, user):
    ops = await _seed_ops(api, session, user)
    await api.post("/transactions", headers=idem(), json={
        "kind": "expense", "amount": "1", "currency": "VND",
        "category_id": await cat_id(session, user.id, "Еда"), "description": "=HYPERLINK(1)"})
    await api.delete(f"/transactions/{ops[1]['id']}")
    r = await api.get("/export/transactions", params={"format": "csv"})
    assert r.status_code == 200
    assert r.headers["content-disposition"].startswith(
        'attachment; filename="pocket-transactions-')
    raw = r.content
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][:6] == ["ID", "Дата", "Время", "Тип", "Сумма", "Валюта"]
    assert len(rows) == 1 + 4  # deleted one excluded
    usd = next(x for x in rows[1:] if x[5] == "USD")
    assert usd[4] == "180" and usd[3] == "Расход" and usd[15] == "vast.ai"
    assert any(x[17] == "'=HYPERLINK(1)" for x in rows[1:])  # formula injection guard
    ton = next(x for x in rows[1:] if x[5] == "TON")
    assert ton[18] == "16.2"  # decimal point
    filtered = await api.get("/export/transactions", params={"format": "csv", "currency": "TON"})
    assert len(list(csv.reader(io.StringIO(filtered.content.decode("utf-8-sig")),
                               delimiter=";"))) == 2

    from openpyxl import load_workbook

    x = await api.get("/export/transactions", params={"format": "xlsx"})
    assert x.status_code == 200 and x.headers["content-type"].startswith(
        "application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb.active
    values = list(ws.iter_rows(values_only=True))
    assert values[0][0] == "ID" and len(values) == 5
    amounts = [row[4] for row in values[1:]]
    assert all(isinstance(a, (int, float)) for a in amounts)  # numeric cells

    j = await api.get("/export/full")
    assert j.status_code == 200 and "pocket-export-" in j.headers["content-disposition"]
    data = j.json()
    assert len(data["transactions"]) == 5  # includes deleted
    assert any(t["deleted_at"] for t in data["transactions"])
    assert {a["name"] for a in data["accounts"]} >= {"Карта", "Кошелёк"}
    kb = next(a for a in data["accounts"] if a["name"] == "Кошелёк")
    assert [x["asset"] for x in kb["assets"]] == ["TON"]
    assert data["settings"]["timezone"] == "Asia/Ho_Chi_Minh"
    assert all(isinstance(t["amount"], str) for t in data["transactions"])


async def test_wipe_and_delete_user(api, session, user, principal):
    await _seed_ops(api, session, user)
    await api.post("/my-accounts", json={"handle": "@me2"})
    bank = await acc_id(session, user.id, "Карта")
    await api.put(f"/accounts/{bank}/opening", json={"asset": "VND", "amount": "1"})
    counts = await wipe_financial_data(session, user.id)
    assert counts["transactions"] == 4 and counts["my_accounts"] == 1
    assert await session.scalar(select(func.count()).where(
        Transaction.user_id == user.id)) == 0
    assert await session.scalar(select(func.count()).where(Posting.user_id == user.id)) == 0
    names = list(await session.scalars(select(Account.name).where(Account.user_id == user.id)
                                       .order_by(Account.sort_order)))
    assert names == ["Карта", "Наличные", "Кошелёк", "Биржа"]  # wipe re-seeds the defaults only
    assert await session.scalar(select(func.count()).where(Category.user_id == user.id)) == 23
    bal = (await api.get("/accounts/balances")).json()
    assert bal["total_vnd"] is None
    # the user + everything with ON DELETE CASCADE (NO ACTION FKs between financial tables)
    await _seed_ops(api, session, user)
    uid = user.id
    await delete_user(session, uid)
    assert await session.scalar(select(func.count()).where(User.id == uid)) == 0
    assert await session.scalar(select(func.count()).where(Transaction.user_id == uid)) == 0
    assert await session.scalar(select(func.count()).where(Account.user_id == uid)) == 0


async def test_unknown_ids_are_404_or_422(api, session, user):
    assert (await api.get("/transactions/999999")).status_code == 404
    assert (await api.patch("/transactions/999999", json={"version": 1, "amount": "1"})
            ).status_code == 404
    assert (await api.get("/accounts/999999")).status_code == 404
    food = await cat_id(session, user.id, "Еда")
    r = await api.post("/transactions", headers=idem(), json={
        "kind": "expense", "amount": "1", "currency": "VND", "category_id": food,
        "account_id": 999999})
    assert r.status_code == 422
    r = await api.post("/transactions", headers=idem(), json={
        "kind": "expense", "amount": "1", "currency": "VND", "category_id": 999999})
    assert r.status_code == 422
    assert (await api.post(f"/actions/{uuid.uuid4()}/undo", headers=idem())).status_code == 404
