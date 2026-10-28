"""Ledger guards: account delete vs a parallel operation,
refund limits, reminders on debt accounts, control characters in XLSX, LIKE escaping in the
merchant search."""

from __future__ import annotations

import asyncio
import io

from sqlalchemy import func, select

from app.models import Account, Transaction
from tests.core.conftest import acc_id, cat_id, idem


def _code(r) -> str:  # type: ignore[no-untyped-def]
    return r.json()["error"]["code"]


# -------------------------------------------------------------------------------


async def test_delete_account_vs_parallel_operation_never_orphans(live_api, live_session,
                                                                  live_user):
    """DELETE /accounts/{id} racing POST /transactions on it: either the operation is first
    (the account is then in use → archived) or the delete is first (→ 422 «Счёт не найден»).
    Never an operation on a soft-deleted, invisible account."""
    async with live_session() as s:
        food = await cat_id(s, live_user.id, "Еда")
    outcomes = set()
    for i in range(20):
        r = await live_api.post("/accounts", json={"name": f"Race {i}", "kind": "bank",
                                                   "currency": "VND"})
        assert r.status_code == 201, r.text
        aid = r.json()["id"]
        d, t = await asyncio.gather(
            live_api.delete(f"/accounts/{aid}"),
            live_api.post("/transactions", headers=idem(), json={
                "kind": "expense", "amount": "1000", "currency": "VND",
                "category_id": food, "account_id": aid}),
        )
        assert d.status_code == 200, d.text
        assert t.status_code in (201, 422), t.text
        outcomes.add(t.status_code)
    async with live_session() as s:
        orphans = await s.scalar(
            select(func.count()).select_from(Transaction).join(
                Account, Account.id == Transaction.account_id
            ).where(Account.user_id == live_user.id, Account.deleted_at.isnot(None))
        )
    assert orphans == 0, outcomes


# -------------------------------------------------------------------------------


async def _expense(api, session, user, amount: str = "100000") -> dict:  # type: ignore[no-untyped-def]
    food = await cat_id(session, user.id, "Еда")
    r = await api.post("/transactions", headers=idem(), json={
        "kind": "expense", "amount": amount, "currency": "VND", "category_id": food})
    assert r.status_code == 201, r.text
    return r.json()["transaction"]


async def _refund(api, session, user, orig_id: int, amount: str, currency: str = "VND"):  # type: ignore[no-untyped-def]
    food = await cat_id(session, user.id, "Еда")
    return await api.post("/transactions", headers=idem(), json={
        "kind": "refund", "amount": amount, "currency": currency, "category_id": food,
        "refund_of_id": orig_id})


async def test_refund_cannot_exceed_the_expense(api, session, user):
    e = await _expense(api, session, user)
    r = await _refund(api, session, user, e["id"], "150000")
    assert r.status_code == 422 and r.json()["error"]["details"]["max"] == "100000"
    r = await _refund(api, session, user, e["id"], "60000")
    assert r.status_code == 201, r.text
    first = r.json()["transaction"]
    r = await _refund(api, session, user, e["id"], "50000")  # 60k + 50k > 100k
    assert r.status_code == 422 and r.json()["error"]["details"]["max"] == "40000"
    assert (await _refund(api, session, user, e["id"], "40000")).status_code == 201
    # PATCH of a refund is checked too (its own amount does not count twice)
    r = await api.patch(f"/transactions/{first['id']}",
                        json={"version": first["version"], "amount": "70000"})
    assert r.status_code == 422, r.text
    r = await api.patch(f"/transactions/{first['id']}",
                        json={"version": first["version"], "amount": "50000"})
    assert r.status_code == 200, r.text
    # another currency is not comparable without a rate: not blocked
    assert (await _refund(api, session, user, e["id"], "5", "USD")).status_code == 201


async def test_refund_of_deleted_or_non_expense_is_422(api, session, user):
    e = await _expense(api, session, user)
    r = await api.delete(f"/transactions/{e['id']}", headers=idem())
    assert r.status_code == 200, r.text
    r = await _refund(api, session, user, e["id"], "1000")
    assert r.status_code == 422 and r.json()["error"]["message"] == "Исходный расход удалён"
    salary = await cat_id(session, user.id, "Зарплата", "income")
    r = await api.post("/transactions", headers=idem(), json={
        "kind": "income", "amount": "1000", "currency": "VND", "category_id": salary})
    inc = r.json()["transaction"]
    r = await _refund(api, session, user, inc["id"], "1000")
    assert r.status_code == 422 and r.json()["error"]["message"] == "Исходный расход не найден"


# -------------------------------------------------------------------------------


async def test_recurring_on_a_debt_account_is_422(api, session, user):
    r = await api.post("/debts/operations", headers=idem(), json={
        "op": "lend", "person_name": "Коля", "amount": "1000", "currency": "VND"})
    assert r.status_code == 201, r.text
    debt_acc = r.json()["debt"]["account_id"]
    base = {"title": "Аренда", "amount": "100", "next_due_on": "2026-11-01"}
    r = await api.post("/recurring", json={**base, "account_id": debt_acc})
    assert r.status_code == 422 and r.json()["error"]["details"]["field"] == "account_id"
    cash = await acc_id(session, user.id, "Наличные")
    r = await api.post("/recurring", json={**base, "account_id": cash})
    assert r.status_code == 201, r.text
    rec = r.json()
    r = await api.patch(f"/recurring/{rec['id']}",
                        json={"version": rec["version"], "account_id": debt_acc})
    assert r.status_code == 422, r.text
    # a deleted account is not found either
    r = await api.post("/accounts", json={"name": "Temp", "kind": "bank", "currency": "VND"})
    tmp = r.json()["id"]
    assert (await api.delete(f"/accounts/{tmp}")).status_code == 200
    r = await api.post("/recurring", json={**base, "account_id": tmp})
    assert r.status_code == 422 and r.json()["error"]["message"] == "Счёт не найден"


# -------------------------------------------------------------------------------


async def test_xlsx_export_with_control_characters(api, session, user):
    from openpyxl import load_workbook

    food = await cat_id(session, user.id, "Еда")
    r = await api.post("/transactions", headers=idem(), json={
        "kind": "expense", "amount": "1000", "currency": "VND", "category_id": food,
        "description": "bell\x07 esc\x1b ok", "counterparty": "tab\tok"})
    assert r.status_code == 201, r.text
    x = await api.get("/export/transactions", params={"format": "xlsx"})
    assert x.status_code == 200, x.text
    ws = load_workbook(io.BytesIO(x.content)).active
    cells = [c for row in ws.iter_rows(values_only=True) for c in row if isinstance(c, str)]
    assert "bell esc ok" in cells
    assert not any("\x07" in c or "\x1b" in c for c in cells)
    c = await api.get("/export/transactions", params={"format": "csv"})
    assert c.status_code == 200


# -------------------------------------------------------------------------------


async def test_merchant_search_escapes_like_and_matches_aliases_case_insensitively(
    api, session, user
):
    from app.models import Merchant

    # an alias stored with capitals (written before aliases were normalised / by hand)
    session.add(Merchant(user_id=user.id, name="Legacy", aliases=["OldAlias"], hits=0))
    await session.flush()
    for body in ({"name": "Back\\slash", "aliases": []},
                 {"name": "Percent 100%", "aliases": []},
                 {"name": "Under_score", "aliases": []},
                 {"name": "Plain", "aliases": ["MixedCase"]}):
        r = await api.post("/merchants", json=body)
        assert r.status_code == 201, r.text

    async def names(q: str) -> set[str]:
        r = await api.get("/merchants", params={"q": q})
        assert r.status_code == 200, r.text
        return {m["name"] for m in r.json()}

    assert await names("\\") == {"Back\\slash"}
    assert await names("%") == {"Percent 100%"}
    assert await names("_") == {"Under_score"}
    assert await names("mixedcase") == {"Plain"}
    assert await names("MIXED") == {"Plain"}
    assert await names("oldalias") == {"Legacy"}


async def test_stale_merchant_default_account_is_not_used(live_session, live_user,
                                                          live_principal):
    """The AI path loads accounts into the session before the ledger takes its lock: a
    merchant's default account deleted meanwhile must be re-read, not taken from the
    identity map."""
    from app.models import Merchant
    from app.schemas.transactions import TransactionCreate
    from app.services.accounts import service as accounts
    from app.services.ledger import transactions as ledger

    async with live_session() as s:
        acc = Account(user_id=live_user.id, name="Old card", kind="card", currency="USD",
                      aliases=[], default_for_currencies=[], sort_order=90)
        s.add(acc)
        await s.flush()
        m = Merchant(user_id=live_user.id, name="Netflix", aliases=[], hits=1,
                     default_account_id=acc.id)
        s.add(m)
        await s.commit()
        acc_id_, m_id = acc.id, m.id
    async with live_session() as s1:
        held = await s1.get(Account, acc_id_)  # in the identity map, not deleted yet
        assert held is not None and held.deleted_at is None
        async with live_session() as s2:
            await accounts.delete_account(s2, live_principal, acc_id_)
            await s2.commit()
        food = await cat_id(s1, live_user.id, "Еда")
        tx = await ledger.create_transaction(s1, live_principal, TransactionCreate(
            kind="expense", amount="15", currency="USD", category_id=food, merchant_id=m_id))
        assert tx.account_id != acc_id_ and held.deleted_at is not None
        await s1.rollback()
